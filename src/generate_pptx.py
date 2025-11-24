from pptx import Presentation

# Create presentation
prs = Presentation()

# Slide content based on user's request
slides_content = [
    ("Titlu", ["NekoNews - Detector de Stiri Anime", "Donea Eduard", "Program MPI", "2025"]),
    ("Motivația proiectului", [
        "Explozia fenomenului de fake news în comunitatea anime",
        "Lipsa verificării credibilității informațiilor",
        "Diferențiere importantă între rumor și fake",
        "Utilizatorii distribuie informații neverificate"
    ]),
    ("Obiectivul proiectului", [
        "Clasificarea automată a știrilor în real, rumor și fake",
        "Generarea unui scor de încredere",
        "Afișarea știrilor similare pentru validare contextuală",
        "Interfață intuitivă pentru utilizator"
    ]),
    ("Colectarea datelor - REAL", [
        "Scraping din AnimeNewsNetwork",
        "Surse alternative RSS",
        "Filtrare, curățare și structurare date",
        "Salvare în format CSV"
    ]),
    ("Colectarea datelor - RUMOR", [
        "Scraping Reddit (r/anime + subreddite dedicate)",
        "Capturarea limbajului informal",
        "Clasificare corectă ca rumor, nu fake"
    ]),
    ("Colectarea datelor - FAKE", [
        "Prima încercare: generare automată (rezultate neconvingătoare)",
        "Soluția adoptată: scraping pentru fake real",
        "Completare cu synthetic data acolo unde a fost util"
    ]),
    ("Construirea datasetului final", [
        "Combinarea surselor într-un singur CSV",
        "Standardizare și verificare dubluri",
        "Distribuția datasetului: rumor > real > fake"
    ]),
    ("Feature Engineering", [
        "TF-IDF cu n-gram (1-2)",
        "max_features = 50,000",
        "Stopwords removal",
        "Feature suplimentar: fake_signal pe cuvinte manipulative"
    ]),
    ("Alegerea modelului", [
        "Model ales: Logistic Regression multiclass",
        "Avantaje: rapid, interpretable, stabil pe texte scurte, fără GPU",
        "Alternative: Naive Bayes, SVM",
        "Future upgrade: BERT / Transformers"
    ]),
    ("Antrenare și Oversampling", [
        "Împărțire stratificată train/test",
        "RandomOverSampler pentru balansare clase",
        "max_iter = 2000",
        "Reproductibilitate (random_state=42)"
    ]),
    ("Evaluarea modelului", [
        "Classification report",
        "Confusion matrix",
        "Performanță ridicată pe real",
        "Confuzii moderate rumor ↔ fake",
        "Posibil overfitting datorită tiparelor evidente"
    ]),
    ("Aplicatia Web", [
        "Framework: Flask",
        "Pagini: Home, Search, Browse",
        "Clasificare live + scor + similar news",
        "Interfață modernă stil anime neon"
    ]),
    ("Concluzii", [
        "Proiect complet funcțional end-to-end",
        "Model performant cu evaluare solidă",
        "Date reale obținute prin scraping",
        "Interfață intuitivă pentru utilizator"
    ]),
    ("Îmbunătățiri viitoare", [
        "BERT fine-tuning",
        "Explainability (LIME/SHAP)",
        "Suport multi-limbaj",
        "Sentiment & stance detection",
        "Extinderea datasetului"
    ])
]

# Create slides
from pptx.shapes.base import BaseShape

for title, bullets in slides_content:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    if slide.shapes.title:
        slide.shapes.title.text = title
    body_shape = slide.placeholders[1]  # type: ignore[assignment]
    if body_shape.has_text_frame:  # type: ignore[union-attr]
        tf = body_shape.text_frame  # type: ignore[union-attr]
        tf.clear()  # Clear existing text
        for bullet in bullets:
            p = tf.add_paragraph()
            p.text = bullet
            p.level = 0

# Save presentation
file_path = "Prezentare_TCRI_NekoNews.pptx"
prs.save(file_path)

print(f"Presentation saved to: {file_path}")
